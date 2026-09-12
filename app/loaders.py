"""
Turn documents of many formats into plain text for the search index.

Supported: PDF, Word (.docx), Excel (.xlsx, .xlsm, .xls), PowerPoint (.pptx),
CSV and TSV, plain text and Markdown, HTML, and JSON / JSON Lines.

Tables - Excel sheets, CSV files, and tables inside Word or PowerPoint - are
written one row per line as "column: value" pairs. The index cuts text into
small chunks, and a bare row of numbers means nothing once it is separated
from its header row. Repeating the column names on every row keeps each chunk
readable on its own, which is what lets a question like "what is Alice's
salary?" find the right row.

Formats that cannot be read directly (old binary .doc and .ppt, images, and
so on) raise UnsupportedFormatError with a hint on how to convert them.

Each reader imports its library only when a file of that type turns up, so a
missing package only affects that one format.
"""

import csv
import datetime
import importlib
import io
import json
import os
from html.parser import HTMLParser
from typing import Any, Callable, Dict, Iterable, List


class UnsupportedFormatError(ValueError):
    """The file type cannot be indexed. The message says how to convert it."""


class MissingDependencyError(ImportError):
    """A reader needs a package that is not installed."""


def _require(module: str, package: str):
    try:
        return importlib.import_module(module)
    except ImportError as exc:
        raise MissingDependencyError(
            f"reading this file type needs the '{package}' package "
            f"(pip install {package})"
        ) from exc


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------


def _clean(value: Any) -> str:
    """One table cell as a single line of text."""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime) and value.time() == datetime.time(0):
        value = value.date()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def rows_to_text(rows: Iterable[Iterable[Any]], label: str = "") -> str:
    """Render table rows as 'column: value' lines, using the first row as header."""
    cleaned = [[_clean(cell) for cell in row] for row in rows]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return ""

    header = cleaned[0]
    prefix = f"{label} | " if label else ""
    if len(cleaned) == 1:
        return prefix + "; ".join(cell for cell in header if cell)

    def column(i: int) -> str:
        return header[i] if i < len(header) and header[i] else f"Column {i + 1}"

    lines = []
    for row in cleaned[1:]:
        pairs = [f"{column(i)}: {value}" for i, value in enumerate(row) if value]
        if pairs:
            lines.append(prefix + "; ".join(pairs))
    return "\n".join(lines)


def _read_text_with_fallback(path: str) -> str:
    """Read a text file, tolerating the non-UTF-8 encodings Excel and Notepad produce."""
    with open(path, "rb") as handle:
        raw = handle.read()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


# ----------------------------------------------------------------------
# Readers, one per format
# ----------------------------------------------------------------------


def read_text(path: str) -> str:
    return _read_text_with_fallback(path)


def read_pdf(path: str) -> str:
    pypdf = _require("pypdf", "pypdf")
    reader = pypdf.PdfReader(path)
    if reader.is_encrypted:
        # Many "protected" PDFs only restrict printing and open with an empty
        # password. A real password cannot be guessed, so say so.
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = 0
        if not unlocked:
            raise UnsupportedFormatError(
                "the PDF is password-protected. Save an unprotected copy first"
            )
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page)


def read_docx(path: str) -> str:
    docx = _require("docx", "python-docx")
    document = docx.Document(path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for number, table in enumerate(document.tables, 1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        text = rows_to_text(rows, f"Table {number}")
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def read_xlsx(path: str) -> str:
    openpyxl = _require("openpyxl", "openpyxl")
    # data_only=True returns the values formulas last calculated to, which is
    # what a person reading the sheet sees, rather than the formula text.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        parts = []
        for sheet in workbook.worksheets:
            text = rows_to_text(sheet.iter_rows(values_only=True), f"Sheet: {sheet.title}")
            if text:
                parts.append(text)
        return "\n\n".join(parts)
    finally:
        workbook.close()


def read_xls(path: str) -> str:
    xlrd = _require("xlrd", "xlrd")
    book = xlrd.open_workbook(path)

    def row_values(sheet, index):
        values = []
        for cell in sheet.row(index):
            if cell.ctype == xlrd.XL_CELL_DATE:
                # .xls stores dates as day counts, so convert them back.
                try:
                    values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    continue
                except Exception:
                    pass
            values.append(cell.value)
        return values

    parts = []
    for sheet in book.sheets():
        rows = (row_values(sheet, i) for i in range(sheet.nrows))
        text = rows_to_text(rows, f"Sheet: {sheet.name}")
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def _read_delimited(path: str, default_delimiter: str) -> str:
    text = _read_text_with_fallback(path)
    try:
        # Excel in many European locales exports CSV with ';', so detect it.
        delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = default_delimiter
    return rows_to_text(csv.reader(io.StringIO(text), delimiter=delimiter))


def read_csv(path: str) -> str:
    return _read_delimited(path, ",")


def read_tsv(path: str) -> str:
    return _read_delimited(path, "\t")


def read_pptx(path: str) -> str:
    pptx = _require("pptx", "python-pptx")
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_text(shape) -> List[str]:
        try:
            shape_type = shape.shape_type
        except Exception:
            shape_type = None
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            out: List[str] = []
            for inner in shape.shapes:
                out.extend(shape_text(inner))
            return out

        out = []
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text:
                out.append(text)
        if getattr(shape, "has_table", False):
            rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
            text = rows_to_text(rows)
            if text:
                out.append(text)
        return out

    presentation = pptx.Presentation(path)
    parts = []
    for number, slide in enumerate(presentation.slides, 1):
        texts: List[str] = []
        for shape in slide.shapes:
            texts.extend(shape_text(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Speaker notes: {notes}")
        if texts:
            parts.append(f"Slide {number}\n" + "\n".join(texts))
    return "\n\n".join(parts)


class _HTMLText(HTMLParser):
    """Collect visible text, with line breaks where block elements end."""

    SKIP = {"script", "style", "noscript", "template", "svg"}
    BLOCK = {
        "p", "div", "br", "li", "tr", "td", "th", "table", "ul", "ol", "pre",
        "section", "article", "header", "footer", "nav", "blockquote", "title",
        "h1", "h2", "h3", "h4", "h5", "h6",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: List[str] = []
        self._skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skipping += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skipping = max(0, self._skipping - 1)
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skipping:
            self.parts.append(data)


def read_html(path: str) -> str:
    parser = _HTMLText()
    parser.feed(_read_text_with_fallback(path))
    parser.close()
    lines = (" ".join(line.split()) for line in "".join(parser.parts).splitlines())
    return "\n".join(line for line in lines if line)


def _flatten_json(value: Any, path: str, out: List[str]) -> None:
    if isinstance(value, dict):
        for key, inner in value.items():
            _flatten_json(inner, f"{path}.{key}" if path else str(key), out)
    elif isinstance(value, list):
        for index, inner in enumerate(value):
            _flatten_json(inner, f"{path}[{index}]", out)
    else:
        text = _clean(value)
        if text:
            out.append(f"{path}: {text}" if path else text)


def read_json(path: str) -> str:
    out: List[str] = []
    _flatten_json(json.loads(_read_text_with_fallback(path)), "", out)
    return "\n".join(out)


def read_jsonl(path: str) -> str:
    records = []
    for number, line in enumerate(_read_text_with_fallback(path).splitlines(), 1):
        if line.strip():
            out: List[str] = []
            _flatten_json(json.loads(line), "", out)
            records.append(f"Record {number}: " + "; ".join(out))
    return "\n".join(records)


# ----------------------------------------------------------------------
# Dispatch
# ----------------------------------------------------------------------

READERS: Dict[str, Callable[[str], str]] = {
    ".txt": read_text,
    ".md": read_text,
    ".markdown": read_text,
    ".rst": read_text,
    ".log": read_text,
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".xlsx": read_xlsx,
    ".xlsm": read_xlsx,
    ".xls": read_xls,
    ".csv": read_csv,
    ".tsv": read_tsv,
    ".pptx": read_pptx,
    ".html": read_html,
    ".htm": read_html,
    ".json": read_json,
    ".jsonl": read_jsonl,
}

SUPPORTED_EXTENSIONS = tuple(sorted(READERS))

FRIENDLY_FORMATS = "PDF, Word, Excel, PowerPoint, CSV, text, Markdown, HTML and JSON"

_CONVERT_HINTS = {
    ".doc": "old Word format. Open it in Word and Save As .docx",
    ".ppt": "old PowerPoint format. Open it in PowerPoint and Save As .pptx",
    ".rtf": "Rich Text format. Save it as .docx or .pdf",
    ".odt": "OpenDocument text. Save it as .docx",
    ".ods": "OpenDocument spreadsheet. Save it as .xlsx",
    ".odp": "OpenDocument slides. Save it as .pptx",
    ".pages": "Apple Pages. Export it as .docx or .pdf",
    ".numbers": "Apple Numbers. Export it as .xlsx",
    ".key": "Apple Keynote. Export it as .pptx or .pdf",
    ".zip": "a zip archive. Extract it into the folder first",
}
for _ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".heic"):
    _CONVERT_HINTS[_ext] = "an image. Text inside pictures needs OCR, which is not set up"


def read_document(path: str) -> str:
    """Return the text of one document, or raise UnsupportedFormatError."""
    ext = os.path.splitext(path)[1].lower()
    reader = READERS.get(ext)
    if reader is None:
        raise UnsupportedFormatError(
            _CONVERT_HINTS.get(ext) or f"'{ext or 'no extension'}' files are not supported"
        )
    # Normalise Windows (\r\n) and old Mac (\r) line endings. Text is decoded
    # from raw bytes, so without this every line keeps a stray \r and the
    # chunker cannot see paragraph breaks, which it looks for as \n\n.
    return reader(path).replace("\r\n", "\n").replace("\r", "\n")
